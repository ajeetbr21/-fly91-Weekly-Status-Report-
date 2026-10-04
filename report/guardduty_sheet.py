"""
Amazon GuardDuty sheet generator for the AWS Weekly BAU Report.

Displays Amazon GuardDuty findings for the reporting week, including a
severity summary (High / Medium / Low) with colour-coded counts and a
detailed findings table.
"""

from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from report.styles import ReportStyles
from report import screenshots
from utils.logger import get_logger

logger = get_logger("bau_report")


class GuardDutySheet:
    """Generates the Amazon GuardDuty worksheet."""

    SHEET_NAME = "GuardDuty"

    SUMMARY_HEADERS = ["Severity", "Count"]

    SEVERITY_ROWS = ["HIGH", "MEDIUM", "LOW"]

    DETAIL_HEADERS = [
        "Finding Type",
        "Title",
        "Severity",
        "Severity Score",
        "Resource",
        "Region",
        "Count",
        "First Seen",
        "Last Seen",
    ]

    def generate(self, wb, guardduty_data, styles=None):
        """
        Create the Amazon GuardDuty sheet.

        Parameters
        ----------
        wb : openpyxl.Workbook
        guardduty_data : dict
            GuardDuty data from the GuardDutyCollector.
        styles : ReportStyles class (optional)

        Returns
        -------
        openpyxl.worksheet.worksheet.Worksheet
        """
        styles = styles or ReportStyles
        ws = wb.create_sheet(self.SHEET_NAME)
        ws.sheet_properties.tabColor = ReportStyles.DARK_BLUE

        guardduty_data = guardduty_data or {}
        severity_counts = guardduty_data.get("severity_counts", {}) or {}
        findings = guardduty_data.get("findings", []) or []
        date_range = guardduty_data.get("date_range", "N/A")
        collection_status = guardduty_data.get("collection_status", "ok")
        collection_error = guardduty_data.get("collection_error")

        total_cols = len(self.DETAIL_HEADERS)
        row = 1

        # ── Title bar ────────────────────────────────────────────────
        styles.apply_title_bar(ws, row, "Amazon GuardDuty - Findings", total_cols)
        row += 2

        # ── Date-range section header ────────────────────────────────
        styles.apply_section_header(
            ws, row, f"Reporting Period: {date_range}", total_cols
        )
        row += 2

        # ── Collection-failure banner ────────────────────────────────
        # When the collector could not query GuardDuty (e.g. no detector,
        # or the IAM role lacks guardduty:ListFindings/GetFindings), make
        # the failure visible instead of showing an empty table that looks
        # identical to a clean week.
        if collection_status == "error":
            row = self._write_collection_error(
                ws, row, total_cols, collection_error
            )

        # ── Severity summary table ────────────────────────────────────
        styles.apply_section_header(ws, row, "Severity Summary", total_cols)
        row += 1

        styles.apply_header_row(ws, row, self.SUMMARY_HEADERS)
        row += 1

        for i, severity in enumerate(self.SEVERITY_ROWS):
            even = i % 2 == 0
            count = int(severity_counts.get(severity, 0) or 0)
            label_cell = styles.apply_data_cell(
                ws, row, 1, severity.title(), even=even
            )
            count_cell = styles.apply_data_cell(ws, row, 2, count, even=even)
            self._apply_severity_fill(severity, label_cell, count_cell)
            row += 1

        # Total findings row
        total_label = styles.apply_data_cell(ws, row, 1, "TOTAL", even=False)
        total_label.font = ReportStyles.FONT_DATA_BOLD
        total_val = styles.apply_data_cell(
            ws, row, 2,
            int(guardduty_data.get("total_findings", len(findings)) or 0),
            even=False,
        )
        total_val.font = ReportStyles.FONT_DATA_BOLD
        row += 2

        # ── Detailed findings table ───────────────────────────────────
        styles.apply_section_header(ws, row, "Detailed Findings", total_cols)
        row += 1

        detail_header_row = row
        styles.apply_header_row(ws, row, self.DETAIL_HEADERS)
        row += 1

        if not findings:
            ws.merge_cells(f"A{row}:{get_column_letter(total_cols)}{row}")
            if collection_status == "error":
                empty_msg = (
                    "Collection failed - findings unavailable (see warning above)."
                )
            else:
                empty_msg = "No GuardDuty findings for this week (clean)."
            cell = ws.cell(row=row, column=1, value=empty_msg)
            cell.font = ReportStyles.FONT_DATA
            cell.alignment = ReportStyles.ALIGN_CENTER
            row += 1
        else:
            for i, finding in enumerate(findings):
                even = i % 2 == 0
                row_data = [
                    finding.get("type", "-"),
                    finding.get("title", "-"),
                    (finding.get("severity_label", "-") or "-").title(),
                    finding.get("severity_score", "-"),
                    finding.get("resource_type", "-"),
                    finding.get("region", "-"),
                    finding.get("count", "-"),
                    finding.get("first_seen", "-"),
                    finding.get("last_seen", "-"),
                ]
                styles.apply_data_row(ws, row, row_data, even=even)

                # Colour-code the severity cell (column 3)
                sev = (finding.get("severity_label", "") or "").upper()
                sev_cell = ws.cell(row=row, column=3)
                self._apply_severity_fill(sev, sev_cell)
                row += 1

        # ── Auto-fit, freeze, filter ──────────────────────────────────
        styles.auto_fit_columns(ws, min_width=12)
        styles.freeze_and_filter(ws, detail_header_row)

        # ── Screenshot-style console visual (one-week findings) ────────
        row += 1
        self._embed_findings_screenshot(ws, row, findings, date_range, styles)

        return ws

    # ------------------------------------------------------------------
    # Private helpers
    # ------------------------------------------------------------------

    def _write_collection_error(self, ws, row, total_cols, error):
        """Render a distinct red banner when GuardDuty collection failed."""
        ws.merge_cells(f"A{row}:{get_column_letter(total_cols)}{row}")
        msg = (
            "⚠ GuardDuty findings could NOT be collected (collection failed / "
            "access denied). This is NOT a clean week - the result below is "
            "empty because the API call did not succeed. Verify GuardDuty is "
            "enabled and the IAM role has guardduty:ListDetectors, "
            "guardduty:ListFindings and guardduty:GetFindings."
        )
        cell = ws.cell(row=row, column=1, value=msg)
        cell.fill = ReportStyles.FILL_RED
        cell.font = Font(name="Calibri", size=10, bold=True, color=ReportStyles.RED_FG)
        cell.alignment = ReportStyles.ALIGN_LEFT
        ws.row_dimensions[row].height = 42
        row += 1
        if error:
            ws.merge_cells(f"A{row}:{get_column_letter(total_cols)}{row}")
            det = ws.cell(row=row, column=1, value=f"Details: {error}")
            det.font = ReportStyles.FONT_SMALL
            det.alignment = ReportStyles.ALIGN_LEFT
            row += 1
        row += 1
        return row

    def _embed_findings_screenshot(self, ws, row, findings, date_range, styles):
        """
        Render and embed a console-screenshot-style PNG of the GuardDuty
        one-week findings (High / Medium / Low severity).

        Any failure (missing Pillow, drawing error, embed error) is logged
        and swallowed so it never aborts the sheet.
        """
        try:
            total_cols = len(self.DETAIL_HEADERS)
            styles.apply_section_header(
                ws, row, "GuardDuty Console - Findings (Screenshot)", total_cols
            )
            anchor_row = row + 1

            headers = ["Severity", "Finding Type", "Title", "Resource", "Count"]
            img_rows = []
            for finding in findings:
                img_rows.append([
                    (finding.get("severity_label", "-") or "-").upper(),
                    finding.get("type", "-"),
                    finding.get("title", "-"),
                    finding.get("resource_type", "-"),
                    finding.get("count", "-"),
                ])

            title = "Amazon GuardDuty - Findings (One Week)"
            if date_range and date_range != "N/A":
                title = f"{title}  |  {date_range}"

            png = screenshots.render_findings_table(
                title, headers, img_rows, severity_col=0
            )
            if png is None:
                return

            from openpyxl.drawing.image import Image as XLImage

            img = XLImage(png)
            ws.add_image(img, f"A{anchor_row}")
        except Exception as exc:
            logger.warning(
                "Could not embed GuardDuty findings screenshot: %s", exc
            )

    @staticmethod
    def _apply_severity_fill(severity, *cells):
        """Apply a RED/YELLOW/GREEN fill to *cells* based on severity."""
        severity = (severity or "").upper()
        if severity == "HIGH":
            fill = ReportStyles.FILL_RED
            fg = ReportStyles.RED_FG
        elif severity == "MEDIUM":
            fill = ReportStyles.FILL_YELLOW
            fg = ReportStyles.YELLOW_FG
        else:
            fill = ReportStyles.FILL_GREEN
            fg = ReportStyles.GREEN_FG

        for cell in cells:
            cell.fill = fill
            cell.font = Font(name="Calibri", size=10, bold=True, color=fg)
