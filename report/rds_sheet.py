"""
RDS sheet generator for the FLY-91 Weekly Report.
"""

from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from report.styles import ReportStyles


class RDSSheet:
    """Generates the RDS Instances worksheet."""

    SHEET_NAME = "RDS Instances"

    HEADERS_1 = [
        "RDS Name",
        "Down Time",
        "Instance type",
        "Minimum Utilization",
        "Maximum Utilization",
        "Average Utilization",
        "Free Memory",
        "Free Storage",
    ]

    HEADERS_2 = [
        "RDS Name",
        "Network Transmit Throughput (Bytes per second)",
        "Network Receive Throughput (Bytes per second)",
        "Max Database Connection (Count)",
    ]

    @staticmethod
    def _pct_cell(display, numeric):
        """Resolve a utilisation cell.

        Prefers an already-formatted display string (e.g. "3.78%"); falls
        back to the numeric cpu_* value rendered with a trailing "%".
        """
        if display is not None and display != "":
            return display
        if numeric is not None:
            return f"{numeric}%"
        return "-"

    def generate(self, wb, rds_data, styles=None):
        styles = styles or ReportStyles
        ws = wb.create_sheet(self.SHEET_NAME)
        ws.sheet_properties.tabColor = ReportStyles.DARK_BLUE

        rds_data = rds_data or []

        # Table 1: Instance & CPU & Storage Info
        total_cols_1 = len(self.HEADERS_1)
        widths_1 = [30, 15, 20, 20, 20, 20, 25, 25]
        for idx, w in enumerate(widths_1, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = w

        row = 1
        styles.apply_title_bar(ws, row, "RDS Instance Details", total_cols_1)
        row += 2

        if not rds_data:
            ws.merge_cells(f"A{row}:{get_column_letter(total_cols_1)}{row}")
            cell = ws.cell(row=row, column=1, value="No RDS instances available.")
            cell.font = Font(name="Calibri", size=12, bold=True, color=ReportStyles.LIGHT_BLUE)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[row].height = 30
            return ws

        styles.apply_header_row(ws, row, self.HEADERS_1)
        row += 1

        for i, db in enumerate(rds_data):
            even = i % 2 == 0

            # RDS Name: prefer display "name", fall back to "db_identifier".
            name = db.get("name") or db.get("db_identifier") or "-"

            # Down Time: prefer display "down_time", fall back to "No".
            down_time = db.get("down_time") or "No"

            # Instance type: prefer display "instance_type", fall back to
            # the numeric/production "instance_class".
            instance_type = db.get("instance_type") or db.get("instance_class") or "-"

            # Utilisation: prefer the already-formatted display strings
            # (e.g. "3.78%"); fall back to numeric cpu_* with a "%" suffix.
            min_util = self._pct_cell(db.get("min_utilization"), db.get("cpu_min"))
            max_util = self._pct_cell(db.get("max_utilization"), db.get("cpu_max"))
            avg_util = self._pct_cell(db.get("avg_utilization"), db.get("cpu_avg"))

            # Free Memory: prefer display "free_memory"; fall back to the
            # computed "X GB out of Y GB" from free_memory_bytes/total_memory_gb.
            mem_str = db.get("free_memory")
            if not mem_str:
                free_mem = db.get("free_memory_bytes")
                tot_mem = db.get("total_memory_gb")
                if free_mem is not None and tot_mem is not None:
                    free_mem_gb = free_mem / (1024 ** 3)
                    mem_str = f"{free_mem_gb:.2f}GB out of {tot_mem} GB"
                else:
                    mem_str = "-"

            # Free Storage: prefer display "free_storage"; fall back to the
            # computed string from free_storage_bytes/allocated_storage_gb.
            stg_str = db.get("free_storage")
            if not stg_str:
                free_stg = db.get("free_storage_bytes")
                tot_stg = db.get("allocated_storage_gb")
                if free_stg is not None and tot_stg is not None:
                    free_stg_gb = free_stg / (1024 ** 3)
                    stg_str = f"{free_stg_gb:.1f}GB out of {tot_stg} GB"
                else:
                    stg_str = "-"

            row_data = [
                name,
                down_time,
                instance_type,
                min_util,
                max_util,
                avg_util,
                mem_str,
                stg_str,
            ]
            styles.apply_data_row(ws, row, row_data, even=even)
            row += 1

        row += 3

        # Table 2: Bandwidth and Network status
        total_cols_2 = len(self.HEADERS_2)
        widths_2 = [30, 45, 45, 35]
        # We only override widths if they are wider than Table 1, or we just rely on Table 1's widths for cols A-D
        # For columns B, C, D which are now wider, we will adjust them
        ws.column_dimensions['B'].width = max(ws.column_dimensions['B'].width, 45)
        ws.column_dimensions['C'].width = max(ws.column_dimensions['C'].width, 45)
        ws.column_dimensions['D'].width = max(ws.column_dimensions['D'].width, 35)

        # Title for Table 2
        cell = ws.cell(row=row, column=1, value="Bandwidth and Network status")
        cell.font = Font(name="Calibri", size=14, bold=True)
        row += 2

        styles.apply_header_row(ws, row, self.HEADERS_2)
        row += 1

        for i, db in enumerate(rds_data):
            even = i % 2 == 0

            name = db.get("name") or db.get("db_identifier") or "-"

            # Network throughput: prefer display strings
            # (e.g. "817.54 KB"); fall back to numeric *_bytes_sec via
            # format_bytes.
            tx_disp = db.get("network_transmit_throughput")
            if not tx_disp:
                tx = db.get("network_tx_bytes_sec")
                tx_disp = styles.format_bytes(tx) if tx is not None else "-"

            rx_disp = db.get("network_receive_throughput")
            if not rx_disp:
                rx = db.get("network_rx_bytes_sec")
                rx_disp = styles.format_bytes(rx) if rx is not None else "-"

            # Max DB connections: prefer display "max_db_connections"; fall
            # back to numeric "db_connections".
            conn = db.get("max_db_connections")
            if conn is None:
                conn = db.get("db_connections")
            if conn is None:
                conn_str = "-"
            elif isinstance(conn, (int, float)):
                conn_str = f"{conn:,.0f}"
            else:
                conn_str = str(conn)

            row_data = [
                name,
                tx_disp,
                rx_disp,
                conn_str,
            ]
            styles.apply_data_row(ws, row, row_data, even=even)
            row += 1

        # Adjust auto-filter to only apply to the first table (or disable it, let's just let format_table do its thing)
        # We don't apply an auto_filter across both tables easily, so we skip it.
        return ws
