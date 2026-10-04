"""
Cover sheet generator for the AWS Weekly BAU Report.

Produces a professional title page with client details, reporting period,
and branding elements.
"""

from datetime import datetime

from openpyxl.styles import Font, PatternFill, Alignment, Border
from openpyxl.utils import get_column_letter

from report.styles import ReportStyles


class CoverSheet:
    """Generates the Cover / Title page of the report workbook."""

    SHEET_NAME = "Cover"

    def generate(self, wb, config, start_date, end_date):
        """
        Create the cover sheet.

        Parameters
        ----------
        wb : openpyxl.Workbook
            Target workbook.
        config : dict
            Application configuration containing 'prepared_by', etc.
        start_date, end_date : str
            Reporting period boundaries (ISO-format strings or display strings).

        Returns
        -------
        openpyxl.worksheet.worksheet.Worksheet
        """
        ws = wb.active if wb.sheetnames == ["Sheet"] else wb.create_sheet(self.SHEET_NAME)
        ws.title = self.SHEET_NAME
        ws.sheet_properties.tabColor = ReportStyles.DARK_BLUE

        # Column widths – 8 columns for layout
        col_widths = [3, 20, 5, 35, 5, 20, 5, 3]
        for idx, w in enumerate(col_widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = w
        total_cols = len(col_widths)

        # ── Dark-blue banner (rows 1-8) ───────────────────────────────
        for r in range(1, 9):
            ws.row_dimensions[r].height = 25
            for c in range(1, total_cols + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill = ReportStyles.FILL_DARK_BLUE

        # Main title – merged across the banner
        ws.merge_cells(f"B2:G3")
        title_cell = ws.cell(row=2, column=2, value="AWS Weekly BAU Report")
        title_cell.font = ReportStyles.FONT_COVER_TITLE
        title_cell.fill = ReportStyles.FILL_DARK_BLUE
        title_cell.alignment = Alignment(horizontal="center", vertical="center")

        # Subtitle
        ws.merge_cells(f"B5:G5")
        sub_cell = ws.cell(row=5, column=2, value="Infrastructure & Cost Management Review")
        sub_cell.font = ReportStyles.FONT_COVER_SUBTITLE
        sub_cell.fill = ReportStyles.FILL_DARK_BLUE
        sub_cell.alignment = Alignment(horizontal="center", vertical="center")

        # ── White area – details ──────────────────────────────────────
        detail_rows = [
            (11, "Client Name", config.get("client_name", "Insync Analytics")),
            (13, "AWS Account ID", config.get("aws_account_id", "179787470151")),
            (15, "Reporting Period", f"{start_date}  to  {end_date}"),
            (17, "Prepared By", config.get("prepared_by", "Cloud Operations Team")),
            (19, "Generated On", datetime.now().strftime("%B %d, %Y  %I:%M %p")),
        ]

        for r, label, value in detail_rows:
            ws.row_dimensions[r].height = 26

            lbl_cell = ws.cell(row=r, column=2, value=label)
            lbl_cell.font = ReportStyles.FONT_COVER_DETAIL_BOLD
            lbl_cell.alignment = Alignment(horizontal="right", vertical="center")

            # Colon separator
            ws.cell(row=r, column=3, value=":").font = ReportStyles.FONT_COVER_DETAIL_BOLD
            ws.cell(row=r, column=3).alignment = Alignment(horizontal="center", vertical="center")

            ws.merge_cells(f"D{r}:F{r}")
            val_cell = ws.cell(row=r, column=4, value=value)
            val_cell.font = ReportStyles.FONT_COVER_DETAIL
            val_cell.alignment = Alignment(horizontal="left", vertical="center")

        # ── Light-blue footer bar ─────────────────────────────────────
        footer_row = 23
        ws.row_dimensions[footer_row].height = 30
        ws.merge_cells(f"A{footer_row}:{get_column_letter(total_cols)}{footer_row}")
        footer_cell = ws.cell(row=footer_row, column=1, value="CONFIDENTIAL – For Internal Use Only")
        footer_cell.font = Font(name="Calibri", size=10, bold=True, color=ReportStyles.WHITE)
        footer_cell.fill = ReportStyles.FILL_LIGHT_BLUE
        footer_cell.alignment = Alignment(horizontal="center", vertical="center")

        # Print settings
        ws.sheet_view.showGridLines = False

        return ws
