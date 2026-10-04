"""
Excel report generator (orchestrator) for the AWS Weekly BAU Report.

Ties all sheet generators together, creates the workbook with correct
properties, invokes each generator in order, and saves the final file.
"""

import logging
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from report.styles import ReportStyles
from report.cover_sheet import CoverSheet
from report.executive_summary import ExecutiveSummarySheet
from report.cost_sheet import CostSheet
from report.ec2_sheet import EC2Sheet
from report.elb_sheet import ELBSheet
from report.waf_sheet import WAFSheet
from report.s3_sheet import S3Sheet
from report.rds_sheet import RDSSheet
from report.dashboard_sheet import DashboardSheet

logger = logging.getLogger(__name__)


class ExcelGenerator:
    """
    Orchestrates generation of the full AWS Weekly BAU Report workbook.

    Usage
    -----
    >>> gen = ExcelGenerator(config, start_date, end_date)
    >>> path = gen.generate(collected_data, "output/report.xlsx")
    """

    # Default configuration values
    DEFAULT_CONFIG = {
        "client_name": "Insync Analytics",
        "aws_account_id": "179787470151",
        "prepared_by": "Cloud Operations Team",
    }

    def __init__(self, config=None, start_date=None, end_date=None):
        """
        Initialise the generator.

        Parameters
        ----------
        config : dict, optional
            Application config (client_name, prepared_by, etc.).
        start_date, end_date : str, optional
            Reporting period boundaries.
        """
        self.config = {**self.DEFAULT_CONFIG, **(config or {})}
        self.start_date = start_date or "N/A"
        self.end_date = end_date or "N/A"

    def generate(self, collected_data, output_path):
        """
        Build the complete report workbook and save to *output_path*.

        Parameters
        ----------
        collected_data : dict
            Data dictionary produced by the collectors.
        output_path : str or pathlib.Path
            Destination file path for the .xlsx file.

        Returns
        -------
        str
            Absolute path of the saved workbook.
        """
        collected_data = collected_data or {}
        output_path = Path(output_path)

        # Ensure parent directory exists
        output_path.parent.mkdir(parents=True, exist_ok=True)

        logger.info("Creating workbook …")
        wb = Workbook()

        # ── Workbook properties ───────────────────────────────────────
        wb.properties.title = "AWS Weekly BAU Report"
        wb.properties.creator = self.config.get("prepared_by", "Cloud Ops")
        wb.properties.company = self.config.get("client_name", "Insync Analytics")
        wb.properties.description = (
            f"Weekly BAU Report for {self.start_date} to {self.end_date}"
        )
        wb.properties.created = datetime.now()

        # Register named styles once
        ReportStyles.register_named_styles(wb)

        # ── Sheet generators (order matters) ──────────────────────────
        generators = [
            ("Cover", self._gen_cover, {}),
            ("Executive Summary", self._gen_exec_summary, {"data": collected_data}),
            ("Cost Analysis", self._gen_cost, {"data": collected_data}),
            ("EC2 Instances", self._gen_ec2, {"data": collected_data}),
            ("Load Balancers", self._gen_elb, {"data": collected_data}),
            ("WAF", self._gen_waf, {"data": collected_data}),
            ("S3 Storage", self._gen_s3, {"data": collected_data}),
            ("RDS Instances", self._gen_rds, {"data": collected_data}),
            ("Dashboard", self._gen_dashboard, {"data": collected_data}),
        ]

        for name, func, kwargs in generators:
            try:
                logger.info("Generating sheet: %s", name)
                func(wb, **kwargs)
            except Exception as exc:
                logger.error("Failed to generate sheet '%s': %s", name, exc, exc_info=True)
                # Create a fallback error sheet so the report still opens
                self._error_sheet(wb, name, str(exc))

        # Remove the default "Sheet" if it was not reused
        if "Sheet" in wb.sheetnames and len(wb.sheetnames) > 1:
            del wb["Sheet"]

        # ── Save ──────────────────────────────────────────────────────
        wb.save(str(output_path))
        logger.info("Report saved to %s", output_path.resolve())
        return str(output_path.resolve())

    # ------------------------------------------------------------------
    # Individual generator wrappers
    # ------------------------------------------------------------------

    def _gen_cover(self, wb):
        CoverSheet().generate(wb, self.config, self.start_date, self.end_date)

    def _gen_exec_summary(self, wb, data):
        ExecutiveSummarySheet().generate(wb, data)

    def _gen_cost(self, wb, data):
        CostSheet().generate(wb, data.get("cost", {}), ReportStyles)

    def _gen_ec2(self, wb, data):
        EC2Sheet().generate(wb, data.get("ec2", []), ReportStyles)

    def _gen_elb(self, wb, data):
        ELBSheet().generate(wb, data.get("elb", []), ReportStyles)

    def _gen_waf(self, wb, data):
        WAFSheet().generate(wb, data.get("waf", []), ReportStyles)

    def _gen_s3(self, wb, data):
        S3Sheet().generate(wb, data.get("s3", []), ReportStyles)

    def _gen_rds(self, wb, data):
        RDSSheet().generate(wb, data.get("rds", []), ReportStyles)

    def _gen_dashboard(self, wb, data):
        DashboardSheet().generate(wb, data, ReportStyles)

    # ------------------------------------------------------------------
    # Fallback error sheet
    # ------------------------------------------------------------------

    @staticmethod
    def _error_sheet(wb, name, message):
        """Create a minimal error placeholder sheet."""
        safe_name = name[:28] + "…" if len(name) > 30 else name
        try:
            ws = wb.create_sheet(f"{safe_name} (err)")
        except Exception:
            ws = wb.create_sheet("Error")
        ws.cell(row=1, column=1, value=f"Error generating {name}")
        ws.cell(row=2, column=1, value=message)
        ws.cell(row=1, column=1).font = ReportStyles.FONT_DATA_BOLD
