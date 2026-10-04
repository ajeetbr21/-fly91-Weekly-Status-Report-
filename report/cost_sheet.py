"""
Cost sheet generator for the AWS Weekly BAU Report.

Produces a cost summary table, top cost-contributing services breakdown,
and cost change analysis with conditional formatting.
"""

from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from report.styles import ReportStyles


class CostSheet:
    """Generates the Cost Analysis worksheet."""

    SHEET_NAME = "Cost Analysis"

    def generate(self, wb, cost_data, styles=None):
        """
        Create the Cost Analysis sheet.

        Parameters
        ----------
        wb : openpyxl.Workbook
        cost_data : dict
            Cost data from the cost collector.
        styles : ReportStyles class (optional, defaults to ReportStyles)

        Returns
        -------
        openpyxl.worksheet.worksheet.Worksheet
        """
        styles = styles or ReportStyles
        ws = wb.create_sheet(self.SHEET_NAME)
        ws.sheet_properties.tabColor = ReportStyles.DARK_BLUE

        cost_data = cost_data or {}
        total_cols = 6

        # Column widths
        widths = [5, 35, 18, 18, 18, 18]
        for idx, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = w

        row = 1

        # ── Title bar ────────────────────────────────────────────────
        styles.apply_title_bar(ws, row, "Cost Analysis", total_cols)
        row += 2

        # ── Weekly Cost Summary ───────────────────────────────────────
        styles.apply_section_header(ws, row, "Weekly Cost Summary", total_cols)
        row += 1

        summary_headers = ["Metric", "Value"]
        styles.apply_header_row(ws, row, summary_headers, col_start=2)
        row += 1

        current = cost_data.get("current_week_total", 0)
        previous = cost_data.get("previous_week_total", 0)
        diff = cost_data.get("difference", 0)
        pct = cost_data.get("pct_change", 0)
        avg_daily = cost_data.get("avg_daily_cost", 0)

        summary_items = [
            ("Previous Week Cost", f"${previous:,.2f}"),
            ("Current Week Cost", f"${current:,.2f}"),
            ("Weekly Difference", f"${diff:+,.2f}"),
            ("% Change", f"{pct:+.2f}%"),
            ("Avg Daily Cost", f"${avg_daily:,.2f}"),
        ]

        for i, (metric, value) in enumerate(summary_items):
            even = i % 2 == 0
            styles.apply_data_cell(ws, row, 2, metric, even=even, align=ReportStyles.ALIGN_LEFT)
            cell = styles.apply_data_cell(ws, row, 3, value, even=even, align=ReportStyles.ALIGN_RIGHT)

            # Colour the difference / % change rows
            if "Difference" in metric or "% Change" in metric:
                if diff > 0:
                    cell.font = Font(name="Calibri", size=10, bold=True, color=ReportStyles.RED_FG)
                elif diff < 0:
                    cell.font = Font(name="Calibri", size=10, bold=True, color=ReportStyles.GREEN_FG)

            row += 1

        row += 1

        # ── Top Cost-Contributing Services ────────────────────────────
        styles.apply_section_header(ws, row, "Top Cost-Contributing Services", total_cols)
        row += 1

        svc_headers = ["#", "Service", "Cost ($)", "% of Total"]
        styles.apply_header_row(ws, row, svc_headers, col_start=2)
        row += 1

        top_services = cost_data.get("top_services", [])
        # Sort descending by cost
        top_services = sorted(top_services, key=lambda s: s.get("cost", 0), reverse=True)
        total_cost = sum(s.get("cost", 0) for s in top_services) or 1  # avoid div/0

        if not top_services:
            ws.merge_cells(f"B{row}:E{row}")
            cell = ws.cell(row=row, column=2, value="No cost data available.")
            cell.font = ReportStyles.FONT_DATA
            cell.alignment = ReportStyles.ALIGN_CENTER
            row += 1
        else:
            for i, svc in enumerate(top_services[:15]):
                even = i % 2 == 0
                svc_name = svc.get("service", "Unknown")
                svc_cost = svc.get("cost", 0)
                svc_pct = (svc_cost / total_cost) * 100

                styles.apply_data_cell(ws, row, 2, i + 1, even=even, align=ReportStyles.ALIGN_CENTER)
                styles.apply_data_cell(ws, row, 3, svc_name, even=even, align=ReportStyles.ALIGN_LEFT)
                styles.apply_data_cell(ws, row, 4, f"${svc_cost:,.2f}", even=even, align=ReportStyles.ALIGN_RIGHT)
                styles.apply_data_cell(ws, row, 5, f"{svc_pct:.1f}%", even=even, align=ReportStyles.ALIGN_RIGHT)
                row += 1

        row += 1

        # ── Service-Level Cost Changes ────────────────────────────────
        styles.apply_section_header(ws, row, "Service-Level Cost Changes", total_cols)
        row += 1

        change_headers = ["Service", "Cost ($)", "Change Indicator"]
        styles.apply_header_row(ws, row, change_headers, col_start=2)
        row += 1

        service_breakdown = cost_data.get("service_breakdown", top_services)
        service_breakdown = sorted(service_breakdown, key=lambda s: s.get("cost", 0), reverse=True)

        for i, svc in enumerate(service_breakdown[:15]):
            even = i % 2 == 0
            svc_name = svc.get("service", "Unknown")
            svc_cost = svc.get("cost", 0)

            styles.apply_data_cell(ws, row, 2, svc_name, even=even, align=ReportStyles.ALIGN_LEFT)
            styles.apply_data_cell(ws, row, 4, f"${svc_cost:,.2f}", even=even, align=ReportStyles.ALIGN_RIGHT)

            # Change indicator
            change = svc.get("change", svc.get("pct_change", None))
            if change is not None:
                try:
                    change = float(change)
                    if change > 0:
                        indicator = f"↑ +{change:.1f}%"
                        color = ReportStyles.RED_FG
                        fill = ReportStyles.FILL_RED
                    elif change < 0:
                        indicator = f"↓ {change:.1f}%"
                        color = ReportStyles.GREEN_FG
                        fill = ReportStyles.FILL_GREEN
                    else:
                        indicator = "→ No change"
                        color = ReportStyles.BLACK
                        fill = ReportStyles.FILL_LIGHT_GRAY if even else ReportStyles.FILL_WHITE
                except (ValueError, TypeError):
                    indicator = "-"
                    color = ReportStyles.BLACK
                    fill = ReportStyles.FILL_LIGHT_GRAY if even else ReportStyles.FILL_WHITE
            else:
                indicator = "-"
                color = ReportStyles.BLACK
                fill = ReportStyles.FILL_LIGHT_GRAY if even else ReportStyles.FILL_WHITE

            ind_cell = styles.apply_data_cell(ws, row, 3, indicator, even=even, align=ReportStyles.ALIGN_CENTER)
            ind_cell.font = Font(name="Calibri", size=10, bold=True, color=color)
            if change is not None and isinstance(change, (int, float)) and change != 0:
                ind_cell.fill = fill

            row += 1

        # ── Auto-fit & freeze ─────────────────────────────────────────
        styles.auto_fit_columns(ws, min_width=12)
        styles.freeze_and_filter(ws, 4)

        return ws
