"""
ELB (ALB) sheet generator for the AWS Weekly BAU Report.

Displays Application Load Balancer metrics including request counts,
connections, LCUs, redirects, processed bytes, and response times.
"""

from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from report.styles import ReportStyles


class ELBSheet:
    """Generates the Load Balancers (ALB) worksheet."""

    SHEET_NAME = "Load Balancers"

    HEADERS = [
        "ALB Name",
        "Requests",
        "Active connection count",
        "New connection count",
        "Consumed Load Balancer Capacity Units",
        "HTTP redirect count",
        "Processed Bytes",
        "Target Response Time",
    ]

    def generate(self, wb, elb_data, styles=None):
        """
        Create the Load Balancers sheet.

        Parameters
        ----------
        wb : openpyxl.Workbook
        elb_data : list[dict]
            ELB/ALB data from the ELB collector.
        styles : ReportStyles class (optional)

        Returns
        -------
        openpyxl.worksheet.worksheet.Worksheet
        """
        styles = styles or ReportStyles
        ws = wb.create_sheet(self.SHEET_NAME)
        ws.sheet_properties.tabColor = ReportStyles.DARK_BLUE

        elb_data = elb_data or []
        total_cols = len(self.HEADERS)

        row = 1

        # ── Title bar ────────────────────────────────────────────────
        styles.apply_title_bar(ws, row, "Application Load Balancers", total_cols)
        row += 2

        # ── Section header ────────────────────────────────────────────
        styles.apply_section_header(ws, row, "ALB Performance Metrics", total_cols)
        row += 1

        # ── Header row ────────────────────────────────────────────────
        header_row = row
        styles.apply_header_row(ws, row, self.HEADERS)
        row += 1

        # ── Data rows ────────────────────────────────────────────────
        if not elb_data:
            ws.merge_cells(f"A{row}:{get_column_letter(total_cols)}{row}")
            cell = ws.cell(row=row, column=1, value="No ALB data available.")
            cell.font = ReportStyles.FONT_DATA
            cell.alignment = ReportStyles.ALIGN_CENTER
            row += 1
        else:
            for i, alb in enumerate(elb_data):
                even = i % 2 == 0
                row_data = [
                    alb.get("name", "-"),
                    styles.format_count(alb.get("requests")),
                    styles.format_count(alb.get("active_connections")),
                    styles.format_count(alb.get("new_connections")),
                    styles.safe_float(alb.get("consumed_lcus"), decimals=2),
                    styles.format_count(alb.get("http_redirect_count")),
                    styles.format_bytes(alb.get("processed_bytes")),
                    self._format_response_time(alb.get("target_response_time")),
                ]
                styles.apply_data_row(ws, row, row_data, even=even)
                row += 1

        # ── Auto-fit, freeze, filter ──────────────────────────────────
        styles.auto_fit_columns(ws, min_width=14)
        styles.freeze_and_filter(ws, header_row)

        return ws

    @staticmethod
    def _format_response_time(value):
        """Format response time in seconds with appropriate precision."""
        if value is None:
            return "-"
        try:
            value = float(value)
            if value < 0.001:
                return f"{value * 1000:.3f} ms"
            if value < 1:
                return f"{value * 1000:.1f} ms"
            return f"{value:.3f} s"
        except (ValueError, TypeError):
            return "-"
