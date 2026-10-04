"""
S3 sheet generator for the AWS Weekly BAU Report.

Displays S3 bucket storage metrics including total size and object counts
with human-readable formatting (GB/TB, K/M suffixes).
"""

from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from report.styles import ReportStyles


class S3Sheet:
    """Generates the S3 Storage worksheet."""

    SHEET_NAME = "S3 Storage"

    HEADERS = [
        "Bucket Name",
        "Total Bucket Size",
        "Total number of Objects",
    ]

    def generate(self, wb, s3_data, styles=None):
        """
        Create the S3 Storage sheet.

        Parameters
        ----------
        wb : openpyxl.Workbook
        s3_data : list[dict]
            S3 data from the S3 collector.
        styles : ReportStyles class (optional)

        Returns
        -------
        openpyxl.worksheet.worksheet.Worksheet
        """
        styles = styles or ReportStyles
        ws = wb.create_sheet(self.SHEET_NAME)
        ws.sheet_properties.tabColor = ReportStyles.DARK_BLUE

        s3_data = s3_data or []
        total_cols = len(self.HEADERS)

        # Column widths
        widths = [45, 22, 25]
        for idx, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = w

        row = 1

        # ── Title bar ────────────────────────────────────────────────
        styles.apply_title_bar(ws, row, "S3 Storage Summary", total_cols)
        row += 2

        # ── Section header ────────────────────────────────────────────
        styles.apply_section_header(ws, row, "Bucket Details", total_cols)
        row += 1

        # ── Header row ────────────────────────────────────────────────
        header_row = row
        styles.apply_header_row(ws, row, self.HEADERS)
        row += 1

        # ── Data rows ────────────────────────────────────────────────
        if not s3_data:
            ws.merge_cells(f"A{row}:{get_column_letter(total_cols)}{row}")
            cell = ws.cell(row=row, column=1, value="No S3 bucket data available.")
            cell.font = ReportStyles.FONT_DATA
            cell.alignment = ReportStyles.ALIGN_CENTER
            row += 1
        else:
            total_size_bytes = 0
            total_objects = 0

            for i, bucket in enumerate(s3_data):
                even = i % 2 == 0
                name = bucket.get("name", "-")
                size_bytes = bucket.get("total_size_bytes")
                objects = bucket.get("total_objects")

                size_display = styles.format_bytes(size_bytes) if size_bytes is not None else "-"
                obj_display = styles.format_count(objects) if objects is not None else "-"

                row_data = [name, size_display, obj_display]
                styles.apply_data_row(ws, row, row_data, even=even)

                # Accumulate totals
                if size_bytes is not None:
                    total_size_bytes += size_bytes
                if objects is not None:
                    total_objects += objects

                row += 1

            # ── Totals row ────────────────────────────────────────────
            row += 1
            tot_data = [
                "TOTAL",
                styles.format_bytes(total_size_bytes),
                styles.format_count(total_objects),
            ]
            styles.apply_data_row(ws, row, tot_data, even=False)
            for c in range(1, total_cols + 1):
                ws.cell(row=row, column=c).font = ReportStyles.FONT_DATA_BOLD

        # ── Auto-fit, freeze, filter ──────────────────────────────────
        styles.auto_fit_columns(ws, min_width=18)
        styles.freeze_and_filter(ws, header_row)

        return ws
