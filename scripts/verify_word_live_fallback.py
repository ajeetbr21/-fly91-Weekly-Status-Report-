#!/usr/bin/env python3
"""
Standalone verification for the Word report's LIVE single-account fallback.

The mock run and scripts/verify_word_report.py both exercise the fully
populated ``collected_data["word_accounts"]`` path. This script instead feeds
``WordReport`` a payload shaped like REAL collected data (no ``word_accounts``
and no ``config.word_report.accounts``), so it drives:

  * ``WordReport._resolve_accounts`` tier 3 -- synthesising a single account
    from the real ``cost`` dict (current_week_total / previous_week_total /
    avg_daily_cost / difference), and
  * ``WordReport._alarms_from_ec2`` -- turning an EC2 instance whose
    ``cpu_max`` exceeds ``thresholds.cpu_warning`` into a CPU-proxy alarm row.

It then loads the generated .docx back with python-docx and asserts that the
single per-account Summary section renders the synthesised account (via its
"<name> - <account_id>" header line) and that the CPU-proxy alarm row is
present. The report is single-account scoped, so the multi-account "Cost
Summary Difference of All AWS Accounts" fleet table must NOT appear.

Usage:
    python3 scripts/verify_word_live_fallback.py [--keep <out.docx>]

Runs on Python 3.9 (no PEP 604 unions). Requires python-docx.
"""

import argparse
import os
import sys
import tempfile
from datetime import date

# Make the repo root importable when run from anywhere.
_REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _REPO_ROOT not in sys.path:
    sys.path.insert(0, _REPO_ROOT)

from docx import Document

from report.word_report import WordReport


def _build_fixture_config():
    """A config WITHOUT word_report.accounts, forcing the tier-3 fallback."""
    return {
        "client_name": "Fly91 Primary",
        "aws_account_id": "123456789012",
        "thresholds": {"cpu_warning": 70, "cpu_critical": 90},
        "word_report": {
            "client_org": "Aptech Limited",
            "submitter_org": "Operisoft Technologies Pvt Ltd",
            "activity_org": "Operisoft",
            # NOTE: deliberately no "accounts" key -> tier-3 live fallback.
        },
    }


def _build_live_shaped_data():
    """A payload shaped like real collected data (no word_accounts)."""
    return {
        "cost": {
            "current_week_total": 1234.56,
            "previous_week_total": 1500.00,
            "avg_daily_cost": 176.37,
            "difference": -265.44,
        },
        "ec2": [
            {
                "name": "web-prod-01",
                "instance_id": "i-0abc123def4567890",
                "region": "us-east-1",
                "cpu_max": 92.4,  # exceeds cpu_warning (70) -> alarm row
            },
            {
                "name": "idle-box",
                "instance_id": "i-0idlecafe00000001",
                "region": "us-east-1",
                "cpu_max": 12.0,  # below threshold -> NO alarm row
            },
        ],
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--keep", default=None,
                        help="Write the generated .docx to this path (otherwise a temp file is used and removed)")
    args = parser.parse_args()

    start_date = date(2026, 9, 14)
    end_date = date(2026, 9, 20)

    config = _build_fixture_config()
    collected_data = _build_live_shaped_data()

    gen = WordReport(config, start_date, end_date)

    out_path = args.keep
    tmp_handle = None
    if out_path is None:
        tmp_handle = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
        tmp_handle.close()
        out_path = tmp_handle.name

    try:
        gen.generate(collected_data, out_path)
        doc = Document(out_path)

        checks = []

        all_para_text = "\n".join(p.text for p in doc.paragraphs)

        # 1. Single per-account Summary section renders the synthesised account.
        checks.append(("per-account Summary section",
                       any(p.text.strip() == "Summary" for p in doc.paragraphs)))
        checks.append(("synthesised account header (name - id)",
                       "Fly91 Primary - 123456789012" in all_para_text))
        checks.append(("Billing and Cost Overview section",
                       "Billing and Cost Overview" in all_para_text))
        # Single-account scope: NO multi-account fleet table / Total Cost row.
        all_cell_pre = "\n".join(
            cell.text
            for t in doc.tables for row in t.rows for cell in row.cells
        )
        full_pre = all_para_text + "\n" + all_cell_pre
        checks.append(("no multi-account fleet table",
                       "Cost Summary Difference of All AWS Accounts" not in full_pre))
        checks.append(("no multi-account Total Cost row",
                       "Total Cost" not in all_cell_pre))

        # 2. CPU-proxy alarm row appears for the over-threshold instance.
        all_cell_text = "\n".join(
            cell.text
            for t in doc.tables for row in t.rows for cell in row.cells
        )
        checks.append(("CPU-proxy alarm server row present",
                       "web-prod-01" in all_cell_text and "i-0abc123def4567890" in all_cell_text))
        checks.append(("alarm metric row present", "CPU 70%" in all_cell_text))
        checks.append(("peak-cpu trigger note present",
                       "Peak CPU 92.4% during the week" in all_cell_text))
        # Below-threshold instance must NOT appear as an alarm.
        checks.append(("below-threshold instance excluded",
                       "idle-box" not in all_cell_text))

        # 3. Resource Utilization & Alarms heading present (alarms were produced).
        checks.append(("Resource Utilization & Alarms heading",
                       "Resource Utilization & Alarms" in all_para_text))

        failed = [name for name, ok in checks if not ok]
        for name, ok in checks:
            print(("PASS" if ok else "FAIL"), name)

        if failed:
            print("\nFAILED CHECKS: %d" % len(failed), file=sys.stderr)
            return 1

        print("\nALL CHECKS PASSED  (live single-account fallback) tables=%d" % len(doc.tables))
        return 0
    finally:
        if tmp_handle is not None:
            try:
                os.unlink(out_path)
            except OSError:
                pass


if __name__ == "__main__":
    sys.exit(main())
